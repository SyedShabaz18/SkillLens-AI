"""SkillLens AI — authenticated, privacy-scoped career analysis application."""
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import csv, hashlib, hmac, io, json, os, secrets, sqlite3, re

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse, Response
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from analyzer import analyze_resume, find_skills

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.getenv("SKILLLENS_DB", str(BASE_DIR / "skilllens.db")))
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_JD_CHARS = 15000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
# Never ship a known session-signing key. For local use, create an ephemeral key when
# no environment value is supplied; sessions reset on restart. Set SKILLLENS_SECRET_KEY
# to a long random value for stable sessions and any non-local deployment.
SECRET_KEY = os.getenv("SKILLLENS_SECRET_KEY") or secrets.token_urlsafe(48)

app = FastAPI(title="SkillLens AI", description="Private, explainable resume-to-role analysis.", version="3.2.0")
app.add_middleware(SessionMiddleware, secret_key=SECRET_KEY, session_cookie="skilllens_session", same_site="lax", https_only=os.getenv("SKILLLENS_HTTPS_ONLY", "0") == "1", max_age=60 * 60 * 12)
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")


@app.middleware("http")
async def security_headers(request: Request, call_next):
    """Apply safe browser defaults without breaking the local single-page app."""
    response = await call_next(request)
    response.headers.setdefault("X-Content-Type-Options", "nosniff")
    response.headers.setdefault("X-Frame-Options", "DENY")
    response.headers.setdefault("Referrer-Policy", "strict-origin-when-cross-origin")
    response.headers.setdefault("Permissions-Policy", "camera=(), microphone=(), geolocation=()")
    if request.url.path.startswith("/api/"):
        response.headers["Cache-Control"] = "no-store"
        response.headers["Pragma"] = "no-cache"
    return response


def connect_db():
    conn = sqlite3.connect(DB_PATH, timeout=10)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def init_db():
    with connect_db() as conn:
        conn.execute("""CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT, email TEXT NOT NULL UNIQUE COLLATE NOCASE,
            display_name TEXT NOT NULL, password_hash TEXT NOT NULL, created_at TEXT NOT NULL)""")
        conn.execute("""CREATE TABLE IF NOT EXISTS analyses (
            id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
            created_at TEXT NOT NULL, filename TEXT NOT NULL, job_title TEXT NOT NULL, score INTEGER NOT NULL,
            result_json TEXT NOT NULL)""")
        # Safe migration from the supplied pre-authentication project.
        columns = {r[1] for r in conn.execute("PRAGMA table_info(analyses)").fetchall()}
        if "user_id" not in columns:
            conn.execute("ALTER TABLE analyses ADD COLUMN user_id INTEGER REFERENCES users(id)")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_analyses_user_created ON analyses(user_id, id DESC)")

@app.on_event("startup")
def startup(): init_db()

@app.get("/")
def home(): return FileResponse(BASE_DIR / "static" / "index.html")

@app.get("/api/health")
def health(): return {"status":"ok", "app":"SkillLens AI", "version":"3.2.0", "supported_formats":["pdf","docx","txt"]}


def password_hash(password: str, salt: bytes | None = None) -> str:
    salt = salt or secrets.token_bytes(16)
    derived = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 310_000)
    return f"pbkdf2_sha256$310000${salt.hex()}${derived.hex()}"


def verify_password(password: str, encoded: str) -> bool:
    try:
        algorithm, iterations, salt_hex, expected = encoded.split("$")
        if algorithm != "pbkdf2_sha256": return False
        actual = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), bytes.fromhex(salt_hex), int(iterations)).hex()
        return hmac.compare_digest(actual, expected)
    except (ValueError, TypeError): return False


def current_user(request: Request):
    uid = request.session.get("user_id")
    if not uid: raise HTTPException(status_code=401, detail="Please log in to continue.")
    with connect_db() as conn:
        user = conn.execute("SELECT id,email,display_name,created_at FROM users WHERE id=?", (uid,)).fetchone()
    if not user:
        request.session.clear()
        raise HTTPException(status_code=401, detail="Session expired. Please log in again.")
    return dict(user)

class AuthPayload(BaseModel):
    email: str = Field(min_length=5, max_length=254)
    password: str = Field(min_length=8, max_length=128)
    display_name: str = Field(default="", max_length=80)

class RoleComparePayload(BaseModel):
    resume_text: str = Field(min_length=30, max_length=30000)
    roles: list[dict] = Field(min_length=2, max_length=5)

class InterviewPayload(BaseModel):
    role: str = Field(min_length=2, max_length=120)
    level: str = Field(default="mid", pattern="^(entry|mid|senior)$")

@app.post("/api/auth/signup")
def signup(payload: AuthPayload, request: Request):
    email = payload.email.strip().lower()
    # Syntax validation only: verifying that a mailbox exists/belongs to the user
    # requires a confirmation email, which must be configured for deployment.
    email_pattern = re.compile(r"^[a-z0-9!#$%&'*+/=?^_`{|}~-]+(?:\.[a-z0-9!#$%&'*+/=?^_`{|}~-]+)*@(?:[a-z0-9](?:[a-z0-9-]{0,61}[a-z0-9])?\.)+[a-z]{2,63}$", re.IGNORECASE)
    if len(email) > 254 or not email_pattern.fullmatch(email):
        raise HTTPException(400, "Enter a valid email address, such as name@example.com. A complete address with a valid domain is required.")
    if len(payload.password) < 8: raise HTTPException(400, "Password must contain at least 8 characters.")
    name = payload.display_name.strip() or email.split("@", 1)[0]
    created = datetime.now(timezone.utc).isoformat(timespec="seconds")
    try:
        with connect_db() as conn:
            cur = conn.execute("INSERT INTO users(email,display_name,password_hash,created_at) VALUES(?,?,?,?)", (email, name, password_hash(payload.password), created))
            uid = cur.lastrowid
    except sqlite3.IntegrityError: raise HTTPException(409, "An account with this email already exists.")
    request.session.clear(); request.session["user_id"] = uid
    return {"user":{"id":uid,"email":email,"display_name":name,"created_at":created}}

@app.post("/api/auth/login")
def login(payload: AuthPayload, request: Request):
    with connect_db() as conn: user = conn.execute("SELECT * FROM users WHERE email=?", (payload.email.strip().lower(),)).fetchone()
    if not user or not verify_password(payload.password, user["password_hash"]): raise HTTPException(401, "Email or password is incorrect.")
    request.session.clear(); request.session["user_id"] = user["id"]
    return {"user":{"id":user["id"],"email":user["email"],"display_name":user["display_name"],"created_at":user["created_at"]}}

@app.post("/api/auth/logout")
def logout(request: Request):
    request.session.clear(); return {"status":"ok"}

@app.get("/api/auth/me")
def me(request: Request): return {"user":current_user(request)}


def extract_text(filename: str, file_bytes: bytes) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(BytesIO(file_bytes))
            if reader.is_encrypted: raise HTTPException(400, "Password-protected PDFs are not supported.")
            return "\n".join((page.extract_text() or "") for page in reader.pages[:20]).strip()
        if ext == ".docx":
            # DOCX is a ZIP container. Reject suspiciously large expansion ratios before parsing.
            import zipfile
            with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
                entries = archive.infolist()
                expanded_size = sum(item.file_size for item in entries)
                if len(entries) > 2000 or expanded_size > 20 * 1024 * 1024:
                    raise HTTPException(400, "This DOCX contains too much embedded data to process safely.")
                if any(item.compress_size == 0 and item.file_size > 1024 * 1024 for item in entries):
                    raise HTTPException(400, "This DOCX contains a suspicious compressed file entry.")
                if any(item.compress_size and item.file_size / item.compress_size > 250 for item in entries):
                    raise HTTPException(400, "This DOCX has an unsafe compression ratio and was rejected.")
            doc = Document(BytesIO(file_bytes)); pieces = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows: pieces.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(pieces).strip()
        if ext == ".txt": return file_bytes.decode("utf-8-sig", errors="replace").strip()
    except HTTPException: raise
    except Exception as exc: raise HTTPException(400, "Could not read this resume. Try a text-based PDF, DOCX, or UTF-8 TXT file.") from exc
    raise HTTPException(400, "Unsupported file type. Upload PDF, DOCX, or TXT.")

@app.post("/api/analyze")
async def analyze(request: Request, resume: UploadFile = File(...), job_title: str = Form("Target role"), job_description: str = Form(...)):
    user = current_user(request)
    filename = Path(resume.filename or "").name; ext = Path(filename).suffix.lower()
    if not filename or ext not in ALLOWED_EXTENSIONS: raise HTTPException(400, "Upload a resume in PDF, DOCX, or TXT format.")
    if len(job_title.strip()) > 120: raise HTTPException(400, "Job title must be 120 characters or fewer.")
    job_description = job_description.strip()
    if len(job_description) < 40: raise HTTPException(400, "Please enter a job description with at least 40 characters.")
    if len(job_description) > MAX_JD_CHARS: raise HTTPException(400, "Job description is too long (maximum 15,000 characters).")
    file_bytes = await resume.read(MAX_FILE_BYTES + 1)
    if not file_bytes: raise HTTPException(400, "The uploaded resume file is empty.")
    if len(file_bytes) > MAX_FILE_BYTES: raise HTTPException(413, "Resume must be 5 MB or smaller.")
    resume_text = extract_text(filename, file_bytes)
    if len(resume_text) < 30: raise HTTPException(400, "Very little text was extracted. Scanned PDFs need OCR first.")
    result = analyze_resume(resume_text, job_description)
    # Lightweight ATS-style checklist; heuristic only, not an actual vendor ATS score.
    low = resume_text.lower()
    sections = {"Contact information": any(x in low for x in ("@", "phone", "linkedin.com")), "Experience": any(x in low for x in ("experience", "employment", "work history")), "Education": "education" in low, "Skills": "skills" in low, "Projects": "projects" in low, "Readable length": 250 <= len(resume_text) <= 12000}
    result["ats_check"] = {"score": round(sum(sections.values()) / len(sections) * 100), "checks": sections, "warnings": (["Consider adding measurable achievements with numbers."] if not any(c.isdigit() for c in resume_text) else []) + (["Keep the resume in a simple, text-readable layout."] if len(resume_text) < 500 else [])}
    result["job_title"] = job_title.strip() or "Target role"; result["resume_filename"] = filename
    result["notice"] = "Resume content is processed in memory and is not saved. Only your analysis summary is stored in your private account history."
    result["interview_questions"] = interview_questions_for(result["job_title"], "mid")[:5]
    created_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    with connect_db() as conn:
        cur = conn.execute("INSERT INTO analyses(user_id,created_at,filename,job_title,score,result_json) VALUES(?,?,?,?,?,?)", (user["id"], created_at, filename[:180], result["job_title"], result["match_score"], json.dumps(result)))
        conn.commit(); result["analysis_id"] = cur.lastrowid; result["created_at"] = created_at
    return result

@app.get("/api/history")
def history(request: Request):
    user = current_user(request)
    with connect_db() as conn: rows = conn.execute("SELECT id,created_at,filename,job_title,score FROM analyses WHERE user_id=? ORDER BY id DESC LIMIT 50", (user["id"],)).fetchall()
    return {"items":[dict(row) for row in rows]}

def owned_analysis(analysis_id: int, user_id: int):
    with connect_db() as conn: row = conn.execute("SELECT result_json,created_at FROM analyses WHERE id=? AND user_id=?", (analysis_id,user_id)).fetchone()
    if row is None: raise HTTPException(404, "Analysis not found.")
    result = json.loads(row["result_json"]); result["analysis_id"] = analysis_id; result.setdefault("created_at", row["created_at"]); return result

@app.get("/api/history/{analysis_id}")
def get_history_item(analysis_id: int, request: Request): return owned_analysis(analysis_id, current_user(request)["id"])

@app.delete("/api/history/{analysis_id}", status_code=204)
def delete_history_item(analysis_id: int, request: Request):
    user_id = current_user(request)["id"]
    with connect_db() as conn:
        cur = conn.execute("DELETE FROM analyses WHERE id=? AND user_id=?", (analysis_id, user_id))
        conn.commit()
    if cur.rowcount == 0:
        raise HTTPException(404, "Analysis not found.")
    return Response(status_code=204)

@app.get("/api/history/{analysis_id}/report.json")
def export_json(analysis_id: int, request: Request):
    result = owned_analysis(analysis_id, current_user(request)["id"])
    return StreamingResponse(io.BytesIO(json.dumps(result, indent=2).encode()), media_type="application/json", headers={"Content-Disposition":f'attachment; filename="skilllens-analysis-{analysis_id}.json"'})

@app.get("/api/history/{analysis_id}/report.csv")
def export_csv(analysis_id: int, request: Request):
    result = owned_analysis(analysis_id, current_user(request)["id"]); output = io.StringIO(); writer = csv.writer(output)
    writer.writerow(["Metric","Value"])
    for key in ("job_title","resume_filename","match_score","skill_overlap_score","context_similarity_score"): writer.writerow([key,result.get(key,"")])
    writer.writerow([]); writer.writerow(["Matched skill","Importance","Resume evidence"])
    for item in result.get("evidence",[]): writer.writerow([item["skill"],item.get("importance",""),item.get("resume_evidence","")])
    writer.writerow([]); writer.writerow(["Skill gap","Importance","Learning resource","URL"])
    for item in result.get("roadmap",[]): writer.writerow([item["skill"],item.get("importance",""),item.get("resource_title",""),item.get("resource_url","")])
    return StreamingResponse(iter([output.getvalue()]),media_type="text/csv",headers={"Content-Disposition":f'attachment; filename="skilllens-analysis-{analysis_id}.csv"'})

@app.get("/api/history/{analysis_id}/report.pdf")
def export_pdf(analysis_id: int, request: Request):
    result = owned_analysis(analysis_id, current_user(request)["id"]); buff = BytesIO(); pdf = canvas.Canvas(buff, pagesize=letter); width,height = letter; y = height-55
    def line(text, size=10, bold=False):
        nonlocal y
        if y < 55: pdf.showPage(); y=height-55
        pdf.setFont("Helvetica-Bold" if bold else "Helvetica", size); pdf.drawString(48,y,str(text)[:115]); y-=17
    line("SkillLens AI — Career Analysis Report",16,True); line(f"Role: {result.get('job_title','Target role')}",12,True); line(f"Match estimate: {result.get('match_score',0)}% | ATS checklist: {result.get('ats_check',{}).get('score','N/A')}%")
    line("Matched skills",12,True)
    for skill in result.get("matched_skills",[]): line("• "+skill)
    line("Skill gaps and roadmap",12,True)
    for item in result.get("roadmap",[]): line(f"• {item.get('skill')} ({item.get('importance','Mentioned')}) — {item.get('resource_title','')}")
    line("Resume checks",12,True)
    for key,val in result.get("ats_check",{}).get("checks",{}).items(): line(f"• {'PASS' if val else 'REVIEW'}: {key}")
    line("This heuristic report is for preparation only, not a hiring prediction.",9)
    pdf.save(); buff.seek(0)
    return StreamingResponse(buff,media_type="application/pdf",headers={"Content-Disposition":f'attachment; filename="skilllens-analysis-{analysis_id}.pdf"'})

@app.delete("/api/history")
def clear_history(request: Request):
    user=current_user(request)
    with connect_db() as conn: conn.execute("DELETE FROM analyses WHERE user_id=?",(user["id"],)); conn.commit()
    return {"status":"ok","message":"Your account analysis history was cleared."}

@app.get("/api/analytics")
def analytics(request: Request):
    user=current_user(request)
    with connect_db() as conn:
        rows=conn.execute("SELECT score,job_title,created_at FROM analyses WHERE user_id=? ORDER BY id DESC LIMIT 50",(user["id"],)).fetchall()
    scores=[r["score"] for r in rows]
    return {"total_analyses":len(rows),"average_score":round(sum(scores)/len(scores)) if scores else 0,"scores":[{"role":r["job_title"],"score":r["score"],"created_at":r["created_at"]} for r in rows]}

@app.post("/api/interview/questions")
def interview_questions(payload: InterviewPayload, request: Request):
    current_user(request)
    return {"role":payload.role,"level":payload.level,"questions":interview_questions_for(payload.role,payload.level)}

def interview_questions_for(role, level="mid"):
    """Return level-specific practice questions with sample answer guidance.

    Examples are frameworks to personalize, not claims about the user's own experience.
    """
    role_l = role.lower()
    if any(k in role_l for k in ["python", "backend", "software", "developer", "engineer"]):
        family = "software"
        entry_topics = [
            ("Explain how you would build and test a simple REST endpoint.", "I would first clarify the endpoint's purpose and request/response format, validate inputs, keep the business logic separate from routing, return appropriate status codes, and add tests for success, invalid input, and failures."),
            ("How would you find the cause of a bug in code you have not written before?", "I would reproduce the issue, read the error and relevant logs, narrow it down with a minimal test or debugger, check assumptions and recent changes, then fix the root cause and add a regression test."),
            ("What is the difference between a list and a dictionary in Python, and when would you use each?", "I would use a list for an ordered collection accessed by position, and a dictionary for key-value lookups such as mapping a user ID to a user record. I would choose based on the access pattern and data meaning."),
        ]
        mid_topics = [
            ("How would you design a REST API that supports pagination, validation, and consistent errors?", "I would define resource-oriented routes and schemas, validate at the boundary, use stable pagination such as a cursor for changing datasets, return consistent error bodies and status codes, and test edge cases and API contracts."),
            ("Describe how you would investigate a slow endpoint in production.", "I would establish a baseline from latency and error metrics, trace the request across dependencies, inspect query plans and external calls, reproduce with representative data, fix the measured bottleneck, and compare metrics after deployment."),
            ("How do you make database changes safely while an application is live?", "I would use backward-compatible migrations: add the new structure first, deploy code that can handle both states, backfill in batches if needed, verify data, and only remove the old structure in a later release."),
        ]
        senior_topics = [
            ("How would you design a backend service for reliability at significantly higher traffic?", "I would begin with service objectives and workload patterns, identify bottlenecks using observability, then consider stateless scaling, caching, database indexing/partitioning, queues, and back-pressure. I would define failure modes, load-test, and roll out changes incrementally."),
            ("Tell me about a technical trade-off you would make between delivery speed and long-term maintainability.", "I would make the trade-off explicit by comparing user impact, operational risk, reversibility, and future cost. I would isolate any short-term compromise, document why it was chosen, add tests/monitoring, and set a clear trigger or date to revisit it."),
            ("How would you lead a production incident and prevent the same class of incident recurring?", "I would prioritize user impact and mitigation, assign clear incident roles, communicate status, and use telemetry to restore service. Afterward I would run a blameless review, identify systemic causes, and track preventive actions to completion."),
        ]
    elif any(k in role_l for k in ["data", "analyst", "bi", "business intelligence"]):
        family = "data"
        entry_topics = [
            ("How would you check a dataset before using it in an analysis?", "I would check row counts, missing values, duplicates, data types, valid ranges, and unexpected categories. I would compare key totals with a trusted source and document exclusions or assumptions before drawing conclusions."),
            ("Explain the difference between WHERE and HAVING in SQL.", "WHERE filters rows before aggregation, while HAVING filters grouped results after GROUP BY. For example, I would use WHERE to limit records to this year and HAVING to keep only groups whose count exceeds a threshold."),
            ("How would you explain a chart's main finding to a non-technical stakeholder?", "I would lead with the decision-relevant takeaway, state the time period and comparison, use a clear chart with labeled units, explain important caveats, and finish with what action the finding supports."),
        ]
        mid_topics = [
            ("How would you investigate a sudden change in a key business metric?", "I would verify the metric definition and data pipeline first, compare the change by segment and time, check seasonality and recent product or process changes, and test plausible explanations before recommending action."),
            ("How would you optimize a slow SQL query and prove the change helped?", "I would inspect the execution plan, joins, filters, and scanned rows; check indexes and data volume; then change one factor at a time. I would compare runtime and resource use on representative data and confirm the results are unchanged."),
            ("How would you design a dashboard that people can use to make decisions?", "I would start with the decisions and audience, define metrics with owners and grain, prioritize a small number of leading indicators, add useful filters and drill-downs, and validate the dashboard with users."),
        ]
        senior_topics = [
            ("How would you create a trusted metric layer across multiple teams?", "I would align stakeholders on definitions and ownership, create documented reusable models, add freshness and quality tests, track lineage, and introduce a change process so metric changes are reviewed and communicated."),
            ("How would you evaluate whether an experiment's result is reliable enough to act on?", "I would check the hypothesis, randomization, sample ratio, power and minimum detectable effect, guardrail metrics, data quality, and practical significance. I would also review segment effects and avoid treating statistical significance alone as business value."),
            ("How would you prioritize analytics work when several teams need support?", "I would compare expected impact, urgency, confidence, effort, and strategic alignment with stakeholders. I would make trade-offs visible, protect time for foundational data quality, and agree on success measures before committing."),
        ]
    elif any(k in role_l for k in ["design", "ux", "ui", "product designer"]):
        family = "design"
        entry_topics = [
            ("How would you turn a short product brief into a first user flow?", "I would identify the target user and goal, list assumptions and edge cases, sketch the happy path and key alternatives, then validate the flow with a teammate or user before polishing visuals."),
            ("What makes a form easy to use and accessible?", "I would use clear labels, logical grouping, keyboard access, visible focus, helpful error messages, appropriate input types, and sufficient contrast. I would test the form with keyboard navigation and different viewport sizes."),
            ("How would you respond if a user found your design confusing?", "I would ask them to show where they got stuck without leading them, look for patterns across sessions, identify the cause, and test a small revision against the original."),
        ]
        mid_topics = [
            ("How would you balance usability findings with a stakeholder's preferred design?", "I would connect the discussion to user evidence and product goals, show the trade-offs, and propose a prototype or usability test. If a decision went another way, I would document the risk and monitor the outcome."),
            ("How would you measure whether a redesign improved the experience?", "I would define a baseline and target tied to the user task, combine behavioral measures such as completion and error rates with qualitative feedback, and segment results so an average does not hide accessibility or user-group issues."),
            ("How do you maintain consistency across a growing product?", "I would use shared components and design tokens, document usage and accessibility states, partner with engineering on implementation, and review exceptions so the system stays useful rather than becoming a rigid constraint."),
        ]
        senior_topics = [
            ("How would you set a product design strategy when research, business goals, and delivery constraints conflict?", "I would synthesize evidence into user and business outcomes, make constraints explicit, prioritize the highest-risk assumptions, and align teams around principles and measurable outcomes. I would revisit the strategy as evidence changes."),
            ("How would you build a research practice that influences roadmap decisions?", "I would create a continuous research plan tied to product questions, combine qualitative and quantitative signals, share concise findings with decision-makers, and track which decisions changed because of evidence."),
            ("How would you mentor designers while preserving quality and team autonomy?", "I would set clear quality principles, use critiques and regular coaching, delegate meaningful ownership, and tailor feedback to each person's growth goals. I would measure success by stronger decisions and team capability, not by centralizing every review."),
        ]
    else:
        family = "general"
        entry_topics = [
            (f"What have you learned that makes you ready for a {role} role?", "I would connect one relevant project, course, or practical task to the role's requirements, explain my individual contribution, and name one lesson I would apply next time."),
            ("How do you organize work when several tasks are due at once?", "I would confirm deadlines and impact, break work into small deliverables, communicate any conflicts early, and track progress so expectations can be reset before a deadline is missed."),
            ("How do you check the quality of your work?", "I would use a checklist appropriate to the task, verify the result against the original requirement, ask for review where useful, and correct issues before handing it over."),
        ]
        mid_topics = [
            (f"How would you improve a recurring process in a {role} team?", "I would map the current process, gather evidence about delays or errors, identify the root cause, pilot a targeted change, and compare the results against a baseline before scaling it."),
            ("Tell me how you handle an ambiguous request from a stakeholder.", "I would clarify the desired outcome, audience, constraints, and definition of done. I would summarize assumptions, propose a small first step, and check alignment before investing in a full solution."),
            ("How do you balance speed, quality, and stakeholder expectations?", "I would make the trade-offs visible, agree on what must not fail, deliver in stages where possible, and communicate risks and progress early rather than surprising stakeholders at the end."),
        ]
        senior_topics = [
            (f"How would you set priorities and success measures for a {role} function?", "I would align priorities to organizational outcomes, establish a small set of measurable indicators, identify dependencies and risks, and review results regularly with stakeholders to adjust investment."),
            ("Describe how you would lead a high-stakes decision with incomplete information.", "I would separate known facts from assumptions, assess reversibility and downside risk, seek targeted evidence from the people closest to the issue, make the decision owner clear, and define signals that would cause us to revisit it."),
            ("How would you raise standards across a team without slowing delivery?", "I would identify the highest-cost recurring problems, introduce lightweight shared practices, coach people through adoption, and measure whether quality improves without adding unnecessary process."),
        ]

    if level == "entry":
        topics = entry_topics
        behavioral = [
            ("Tell me about a project that shows your readiness for this role.", "Use a project from study, volunteering, or personal work. Explain the goal, your own contribution, the tools you used, one obstacle, and what you learned. Be honest about what was individual versus team work."),
            ("Describe feedback you received and how you used it.", "Choose a specific piece of feedback, explain how you responded without becoming defensive, describe the change you made, and share what improved or what you would do differently next time."),
        ]
        scenario = ("If you did not know how to solve a task, what would you do first?", "I would clarify the expected outcome, break the task into smaller questions, consult reliable documentation or examples, try a small test, and ask a focused question after explaining what I had already tried.")
    elif level == "senior":
        topics = senior_topics
        behavioral = [
            ("How have you helped other people improve their work?", "Describe how you identified a development need, agreed on a goal, provided coaching or useful feedback, and followed up. Focus on the other person's growth and the team's result."),
            ("Describe a decision you owned that had meaningful impact.", "Explain the context and stakes, alternatives considered, evidence and risks, how you aligned people, and the measurable outcome. Include what you learned and what you would change."),
        ]
        scenario = ("A critical project is at risk and teams disagree on the solution. How do you lead?", "I would align everyone on the user or business impact, establish facts and constraints, make options and risks explicit, assign decision ownership, agree on a short-term mitigation, and communicate the decision and follow-up measures.")
    else:
        topics = mid_topics
        behavioral = [
            ("Tell me about a time you improved a process or outcome.", "Describe the baseline problem, your role, the steps you took, how you worked with others, and the result. Use a metric if you have one; otherwise state the observable change without inventing numbers."),
            ("Describe a disagreement with a teammate or stakeholder.", "Explain the different priorities fairly, how you listened and used evidence, how the decision was reached, and what happened afterward. Show collaboration rather than blaming the other person."),
        ]
        scenario = ("A deadline is approaching but you discover a significant quality risk. What do you do?", "I would assess the risk and impact quickly, tell the relevant people early, offer options such as reducing scope or extending the deadline, and recommend the safest workable path. I would not hide the risk to appear on schedule.")

    items = []
    for question, answer in topics:
        items.append({"category": "Technical · " + level.title(), "question": question, "practice_tip": "Structure your response as context → approach → trade-off → result.", "sample_answer": answer})
    for question, answer in behavioral:
        items.append({"category": "Behavioral · " + level.title(), "question": question, "practice_tip": "Use STAR: Situation, Task, Action, Result. Be specific about your own contribution.", "sample_answer": answer})
    items.append({"category": "Scenario · " + level.title(), "question": scenario[0], "practice_tip": "Explain your reasoning, the first step you would take, and how you would judge success.", "sample_answer": scenario[1]})
    # The level changes the technical, behavioral, and scenario prompts. Include a
    # visible level tag and rotate the order per request so repeated practice feels fresh.
    import random
    random.SystemRandom().shuffle(items)
    return items

@app.post("/api/roles/compare")
def compare_roles(payload: RoleComparePayload, request: Request):
    current_user(request); resume_skills={x["skill"] for x in find_skills(payload.resume_text)}; output=[]
    for role in payload.roles:
        title=str(role.get("title","Target role"))[:120]; description=str(role.get("description",""))[:15000]
        if len(description)<20: continue
        analysis=analyze_resume(payload.resume_text,description)
        output.append({"title":title,"match_score":analysis["match_score"],"matched_skills":analysis["matched_skills"],"missing_skills":analysis["missing_skills"],"required_gaps":analysis.get("required_gaps",[])})
    if len(output)<2: raise HTTPException(400,"Provide at least two roles with descriptions of 20+ characters.")
    return {"roles":sorted(output,key=lambda x:x["match_score"],reverse=True),"notice":"Role scores are heuristic comparisons, not hiring probabilities."}
